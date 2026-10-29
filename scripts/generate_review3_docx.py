"""Generate the editable (.docx) Review 3 document from the same content as the PDF.

    DOC_MODEL=/tmp/model.json <venv>/bin/python scripts/generate_review3_pdf.py /tmp/x.pdf
    <venv with python-docx>/bin/python scripts/generate_review3_docx.py /tmp/model.json [output.docx]

The PDF generator exports its content (headings, text panels, tables, figures)
to the JSON model; this script renders that model with real Word headings,
tables and pictures so every line stays editable.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
model = json.loads(Path(sys.argv[1]).read_text())
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "docs/PROJECT_REVIEW_3_IMPLEMENTATION_DOCUMENTATION.docx"

NAVY, BLUE = "17335B", "1F4E79"
FONT = "Arial"
CONTENT_W = 6.5  # inches


def rgb(h):
    return RGBColor.from_string(h)


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def set_cell_bg(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def set_par_bg(par, fill):
    shade(par._p.get_or_add_pPr(), fill)


def borders(table, color="222222", sz=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def run_fmt(r, size=None, bold=None, italic=None, color=None, mono=False):
    r.font.name = "Courier New" if mono else FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New" if mono else FONT)
    if size: r.font.size = Pt(size)
    if bold is not None: r.bold = bold
    if italic is not None: r.italic = italic
    if color: r.font.color.rgb = rgb(color)


TAG = re.compile(r"(<b>|</b>|<i>|</i>|<font name='Mono'>|</font>)")


def add_marked(par, text, size=9, color="333333", bold=False, italic=False):
    """Write text containing the PDF generator's inline tags as formatted runs."""
    b, i, m = bold, italic, False
    for tok in TAG.split(text):
        if tok == "<b>": b = True
        elif tok == "</b>": b = bold
        elif tok == "<i>": i = True
        elif tok == "</i>": i = italic
        elif tok == "<font name='Mono'>": m = True
        elif tok == "</font>": m = False
        elif tok:
            tok = tok.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
            r = par.add_run(tok)
            run_fmt(r, size=size - 0.8 if m else size, bold=b, italic=i, color=color, mono=m)


def spacing(par, before=0, after=4, line=None):
    pf = par.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line: pf.line_spacing = line


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin, sec.bottom_margin = Inches(0.9), Inches(0.8)

normal = doc.styles["Normal"]
normal.font.name = FONT; normal.font.size = Pt(9)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
for name, size, color, before in (("Heading 1", 15, "1B365D", 12), ("Heading 2", 11, BLUE, 8)):
    st = doc.styles[name]
    st.font.name = FONT; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = rgb(color)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT); st.element.rPr.rFonts.set(qn("w:ascii"), FONT); st.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(5); st.paragraph_format.keep_with_next = True

# header / footer as in the template
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run("PROJECT REVIEW 3 | IMPLEMENTATION DOCUMENTATION"); run_fmt(r, 7.6, True, color="718096")
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Review 3 - 100% Implementation Required"); run_fmt(r, 7.6, color="718096")


def fix_widths(table, widths_in):
    """Fixed layout + explicit grid so Word and other viewers keep the intended column widths."""
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for w in tblPr.findall(qn("w:tblW")): tblPr.remove(w)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(widths_in) * 1440))); tw.set(qn("w:type"), "dxa"); tblPr.append(tw)
    for col, w in zip(table.columns, widths_in):
        col.width = Inches(w)
        for c in col.cells: c.width = Inches(w)


def big_heading(text, new_page=True):
    p = doc.add_paragraph(); r = p.add_run(text); run_fmt(r, 15, True, color="1B365D"); spacing(p, 0, 8)
    p.paragraph_format.page_break_before = new_page


def grid_table(header, rows, widths_in, centre=(), bold_first=False, mono_cols=(), hdr_blank_first=False):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    borders(t); cell_margins(t)
    for c, (cell, h) in enumerate(zip(t.rows[0].cells, header)):
        set_cell_bg(cell, NAVY); cell.width = Inches(widths_in[c])
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, 0, 0)
        add_marked(p, h, 8, "FFFFFF", bold=True)
    # repeat header row on each page
    trPr = t.rows[0]._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
    for row in rows:
        cells = t.add_row().cells
        for c, (cell, txt) in enumerate(zip(cells, row)):
            cell.width = Inches(widths_in[c])
            p = cell.paragraphs[0]; spacing(p, 0, 0)
            if c in centre: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if c in mono_cols:
                r = p.add_run(txt.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")); run_fmt(r, 7.4, color="333333", mono=True)
            else:
                add_marked(p, txt, 7.8, "222222" if (bold_first and c == 0) else "444444", bold=bold_first and c == 0)
    fix_widths(t, widths_in)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---------------------------------------------------------------- title page
for _ in range(3): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_fmt(p.add_run("PROJECT REVIEW 3"), 30, True, color="1B365D"); spacing(p, 0, 6)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_fmt(p.add_run("IMPLEMENTATION DOCUMENTATION TEMPLATE"), 17, True, color=BLUE); spacing(p, 0, 16)
band = doc.add_table(rows=1, cols=1); band.alignment = WD_TABLE_ALIGNMENT.CENTER
c = band.rows[0].cells[0]; set_cell_bg(c, "D9EAF5")
p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, 8, 8)
run_fmt(p.add_run("REVIEW 3 REQUIREMENT: 100% IMPLEMENTATION COMPLETED"), 11.5, True, color="1B365D")
fix_widths(band, [CONTENT_W])
doc.add_paragraph()
meta = [("Project Title", "IntelliCannula: Two-Module Clinical Decision Support for Peripheral IV Cannulation (Pre-Cannulation Vein Guidance & Post-Cannulation Infection Monitoring)"),
        ("Student Name(s)", "Sastika M, Akshaya GK"), ("Registration Number(s)", "23MIA1047, 23MIA1127"),
        ("Programme / Branch", "M.Tech. Integrated Computer Science and Engineering (Business Analytics)"),
        ("Course / Project Code", "SWE1017 — Project Review 3 (Final Implementation Review)"),
        ("Faculty Guide", "Dr. Manjula V. (Professor, School of Computer Science and Engineering)"),
        ("Review", "Review 3 - Final Implementation Review")]
mt = doc.add_table(rows=0, cols=2); mt.alignment = WD_TABLE_ALIGNMENT.CENTER; mt.autofit = False
borders(mt, "9DB7CC", 4); cell_margins(mt, 90, 90, 90, 90)
for a, b in meta:
    cells = mt.add_row().cells
    cells[0].width, cells[1].width = Inches(2.0), Inches(4.5)
    set_cell_bg(cells[0], "EAF2F8")
    pa = cells[0].paragraphs[0]; spacing(pa, 0, 0); run_fmt(pa.add_run(a), 9.5, True, color="1B365D")
    pb = cells[1].paragraphs[0]; spacing(pb, 0, 0); run_fmt(pb.add_run(b), 9.5, color="333333")
fix_widths(mt, [2.0, 4.5])

# ---------------------------------------------------------------- checklist
big_heading("Review 3 Submission Checklist")
items = ["Complete source code submitted", "All modules integrated and executable", "Dataset / database / sample input included or access provided",
         "Dependencies / requirements file included", "README / execution instructions included", "Screenshots and outputs included",
         "Testing and validation completed", "Results / performance metrics reported", "Documentation completed"]
grid_table(["", "Item", "Status"], [["☑", i, "Yes"] for i in items], [0.4, 4.7, 1.4], centre=(0,))
p = doc.add_paragraph(); spacing(p, 4, 0)
add_marked(p, "Submission notes: the trained checkpoints (<font name='Mono'>vein_module/models/</font>, <font name='Mono'>infection_module/models/</font>), the raw/processed datasets and the demo images (<font name='Mono'>data/demo/</font>) are git-ignored because of their size, so they are supplied in the submission folder next to the repository. The public datasets can be re-downloaded with <font name='Mono'>scripts/download_datasets.py</font> (Section 13).", 8, "555555", italic=True)

# ---------------------------------------------------------------- contents
big_heading("Contents")
for line in ["1. Problem Statement and Objectives", "2. Finalized System Architecture / Methodology", "3. Module-wise Implementation",
             "4. Algorithms, Models and Technologies Used", "5. Dataset / Database Details", "6. Implementation Screenshots", "7. Experimental Setup",
             "8. Testing and Test Cases", "9. Results and Performance Metrics", "10. Comparison and Analysis", "11. Limitations", "12. Conclusion",
             "13. Instructions for Executing the Project", "Appendix A. Source Code Structure", "Appendix B. Dependencies and Configuration", "Appendix C. Additional Evidence"]:
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35); spacing(p, 0, 5)
    run_fmt(p.add_run(line), 10, color="1B365D")

# ---------------------------------------------------------------- body from the shared model
first_h1 = True
for item in model:
    kind = item[0]
    if kind == "h1":
        h = doc.add_heading(item[1], level=1)
        if first_h1:
            h.paragraph_format.page_break_before = True; first_h1 = False
    elif kind == "h2":
        doc.add_heading(item[1], level=2)
    elif kind == "note":
        p = doc.add_paragraph(); spacing(p, 0, 4); p.paragraph_format.keep_with_next = True
        add_marked(p, item[1], 8, "555555", italic=True)
    elif kind == "shaded":
        for txt in item[1]:
            p = doc.add_paragraph(); spacing(p, 0, 0, 1.1)
            p.paragraph_format.left_indent = Inches(0.05); p.paragraph_format.right_indent = Inches(0.05)
            set_par_bg(p, "F2F2F2")
            add_marked(p, txt, 8.6, "333333", italic=item[2])
            p.paragraph_format.space_after = Pt(3)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    elif kind == "table":
        _, header, rows, widths, centre, mono_cols, bold_first = item
        scale = CONTENT_W / sum(widths)
        grid_table(header, rows, [w * scale for w in widths], tuple(centre), bold_first, tuple(mono_cols))
    elif kind == "steps":
        for n, txt in enumerate(item[1], 1):
            p = doc.add_paragraph(); spacing(p, 0, 0); set_par_bg(p, "F2F2F2"); p.paragraph_format.left_indent = Inches(0.05)
            add_marked(p, f"Step {n}: {txt}", 8, "444444", italic=True)
            p.paragraph_format.space_after = Pt(3)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    elif kind == "figure":
        _, path, caption, max_w_pt, max_h_pt = item
        w, h = Image.open(path).size
        sc = min(max_w_pt / w, max_h_pt / h)
        width_in = min(CONTENT_W, w * sc / 72)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, 4, 2); p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Inches(width_in))
        cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(cp, 0, 8)
        add_marked(cp, caption, 8.2, "333333", italic=True)

# ---------------------------------------------------------------- declaration
doc.add_heading("Student / Team Declaration", level=1)
p = doc.add_paragraph(); spacing(p, 0, 8)
add_marked(p, "We confirm that the implementation documented in this report represents the completed Review 3 version of the project and that the submitted code corresponds to the demonstrated system.", 9.5, "222222")
t = grid_table(["Name", "Registration No.", "Signature", "Date"],
               [["Sastika M", "23MIA1047", "", "29-10-2026"], ["Akshaya GK", "23MIA1127", "", "29-10-2026"]], [2.0, 1.6, 1.6, 1.3])
for row in t.rows[1:]:
    row.height = Inches(0.35)

doc.core_properties.title = "Project Review 3 - Implementation Documentation - IntelliCannula"
doc.core_properties.author = "Sastika M, Akshaya GK"
doc.save(OUT)
print("wrote", OUT)
