import os
import subprocess
import sys

def install_pptx():
    try:
        import pptx
    except ImportError:
        print("Installing python-pptx...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])

install_pptx()

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()

# Slide 1: Title Slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "IntelliCannula"
subtitle.text = "Clinical Decision Support for Peripheral IV Cannulation\nPre-cannulation vein guidance & post-cannulation infection monitoring"

# Slide 2: Introduction
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Introduction"
tf = body.text_frame
tf.text = "Peripheral Intravenous Cannulation (PIVC) is the most common invasive hospital procedure."
p = tf.add_paragraph()
p.text = "High failure rates occur in patients with Difficult Venous Access (DVA)."
p = tf.add_paragraph()
p.text = "Post-cannulation complications (phlebitis, infiltration) affect up to 30% of patients and are often caught late."
p = tf.add_paragraph()
p.text = "IntelliCannula is a two-phase AI decision support system that recommends safe insertion sites and monitors the site over time for early signs of infection."

# Slide 3: Objectives
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Objectives"
tf = body.text_frame
tf.text = "Phase 1 (Pre-Cannulation): Segment subcutaneous veins from Near-Infrared (NIR) imagery and compute optimal insertion coordinates and needle angles."
p = tf.add_paragraph()
p.text = "Phase 2 (Post-Cannulation): Detect cannula hubs accurately and track site erythema (redness) and texture variations over consecutive days."
p = tf.add_paragraph()
p.text = "Clinical Safety: Provide an advisory, human-in-the-loop dashboard following the principle: 'The system recommends, never acts'."

# Slide 4: Overall Methodology (1/2)
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Overall Methodology (1/2): Vein Guidance"
tf = body.text_frame
tf.text = "Dataset: CUBITAL Near-Infrared (NIR) dataset."
p = tf.add_paragraph()
p.text = "Model: Attention U-Net architecture (7.85M parameters) chosen for precise spatial segmentation of venous structures."
p = tf.add_paragraph()
p.text = "Post-Processing Pipeline:"
p2 = tf.add_paragraph()
p2.text = "Skeletonization of the predicted vein mask."
p2.level = 1
p3 = tf.add_paragraph()
p3.text = "Identification of the widest, straightest segment."
p3.level = 1
p4 = tf.add_paragraph()
p4.text = "Calculation of the optimal insertion angle and coordinates."
p4.level = 1

# Slide 5: Overall Methodology (2/2)
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Overall Methodology (2/2): Infection Monitoring"
tf = body.text_frame
tf.text = "Hub Detection (YOLOv11): Trained via domain-randomized synthetic compositing (3D hubs overlaid on real skin backgrounds) to preserve patient privacy while achieving mAP50 of 0.995."
p = tf.add_paragraph()
p.text = "Image Registration: ORB keypoint matching aligns daily RGB clinical photos to the patient's Day 0 baseline to correct camera angles."
p = tf.add_paragraph()
p.text = "Risk Prediction: LightGBM model analyzes temporal color shifts (LAB space) and texture features to output a normalized infection risk score."

# Slide 6: System Architecture
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "System Architecture"
tf = body.text_frame
tf.text = "Frontend: Next.js Clinician Dashboard (React, Tailwind CSS) providing an intuitive interface for image uploads and trend visualization."
p = tf.add_paragraph()
p.text = "Backend: FastAPI (Python) handling REST API requests, audit logging, and ML orchestration."
p = tf.add_paragraph()
p.text = "Database: Resilient architecture using PostgreSQL with an automatic fallback to local SQLite."
p = tf.add_paragraph()
p.text = "ML Integration: Lazy-loaded PyTorch and LightGBM models running in isolated subprocesses to ensure memory efficiency and avoid OpenMP collisions."

# Slide 7: Implementation - Dashboard
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Implementation: Clinician Dashboard"
tf = body.text_frame
tf.text = "[Please insert a screenshot of the IntelliCannula Dashboard Home Page here]"
p = tf.add_paragraph()
p.text = "Features a patient directory, clinical advisory warnings, and pipeline metric badges."

# Slide 8: Output - Pre-Cannulation Vein Scan
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Output: Pre-Cannulation Vein Scan"
tf = body.text_frame
tf.text = "[Please insert a screenshot of the Vein Scan Output here]"
p = tf.add_paragraph()
p.text = "Highlights the vein segmentation overlay on the NIR image."
p = tf.add_paragraph()
p.text = "Displays calculated insertion site coordinates and recommended needle angle."

# Slide 9: Output - Post-Cannulation Infection Check
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Output: Post-Cannulation Infection Check"
tf = body.text_frame
tf.text = "[Please insert a screenshot of the Site Check Output here]"
p = tf.add_paragraph()
p.text = "Shows the localized ROI bounding box detected by YOLOv11."
p = tf.add_paragraph()
p.text = "Displays the temporal risk trend chart and automated clinical alert banners based on risk velocity."

# Slide 10: Conclusion & Future Scope
slide = prs.slides.add_slide(prs.slide_layouts[1])
title, body = slide.shapes.title, slide.placeholders[1]
title.text = "Conclusion & Future Scope"
tf = body.text_frame
tf.text = "Conclusion: IntelliCannula successfully demonstrates end-to-end clinical workflow automation, reducing reliance on subjective visual inspection for PIVC management."
p = tf.add_paragraph()
p.text = "Future Scope:"
p2 = tf.add_paragraph()
p2.text = "Clinical trials with real patient temporal dataset validation."
p2.level = 1
p3 = tf.add_paragraph()
p3.text = "Edge deployment on mobile devices via ONNX runtime for offline bedside inference."
p3.level = 1

prs.save('IntelliCannula_Presentation.pptx')
print("Successfully generated IntelliCannula_Presentation.pptx")
